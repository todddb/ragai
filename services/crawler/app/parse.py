from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Dict, List, Tuple
from urllib.parse import urljoin

try:
    from bs4 import BeautifulSoup  # type: ignore
except Exception:  # pragma: no cover - optional dependency guard
    BeautifulSoup = None  # type: ignore

try:
    import openpyxl  # type: ignore
except Exception:  # pragma: no cover - optional dependency guard
    openpyxl = None  # type: ignore


@dataclass
class ParsedDocument:
    title: str
    markdown: str
    text_for_chunking: str
    links: List[str] = field(default_factory=list)
    meta: Dict = field(default_factory=dict)


def _normalize_content_type(content_type: str) -> str:
    return content_type.split(";")[0].strip().lower()


def _collapse_whitespace(text: str) -> str:
    return " ".join(text.split())


def _decode_text(content_bytes: bytes) -> str:
    return content_bytes.decode("utf-8", errors="replace")


def _parse_html(content_bytes: bytes, base_url: str) -> ParsedDocument:
    if BeautifulSoup is None:
        raise RuntimeError("Missing dependency: beautifulsoup4 (pip install beautifulsoup4)")
    html = _decode_text(content_bytes)
    soup = BeautifulSoup(html, "html.parser")
    title = ""
    if soup.title and soup.title.string:
        title = soup.title.string.strip()
    if not title:
        h1 = soup.find("h1")
        if h1:
            title = h1.get_text(strip=True)
    text = _collapse_whitespace(soup.get_text(separator=" "))
    links = []
    for tag in soup.find_all("a"):
        href = tag.get("href")
        if not href:
            continue
        links.append(urljoin(base_url, href))
    markdown = text
    return ParsedDocument(title=title, markdown=markdown, text_for_chunking=text, links=links)


def _parse_text(content_bytes: bytes) -> ParsedDocument:
    text = _collapse_whitespace(_decode_text(content_bytes))
    return ParsedDocument(title="", markdown=text, text_for_chunking=text)


def _parse_xlsx(content_bytes: bytes) -> ParsedDocument:
    if openpyxl is None:
        raise RuntimeError("Missing dependency: openpyxl (pip install openpyxl)")
    workbook = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
    markdown_lines: List[str] = []
    text_parts: List[str] = []
    for sheet in workbook.worksheets:
        markdown_lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if not any(values):
                continue
            markdown_lines.append(" | ".join(values))
            text_parts.extend(value for value in values if value)
    markdown = "\n".join(markdown_lines).strip()
    text = _collapse_whitespace(" ".join(text_parts))
    return ParsedDocument(
        title="",
        markdown=markdown,
        text_for_chunking=text,
        meta={"sheet_count": len(workbook.worksheets)},
    )


def _parse_binary(content_bytes: bytes) -> ParsedDocument:
    text = _collapse_whitespace(_decode_text(content_bytes))
    return ParsedDocument(title="", markdown=text, text_for_chunking=text)


def parse_by_type(
    content_bytes: bytes, content_type: str, url: str
) -> Tuple[ParsedDocument, str]:
    normalized = _normalize_content_type(content_type or "")
    if normalized in {"text/html", "application/xhtml+xml"}:
        return _parse_html(content_bytes, url), "html"
    if normalized in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    }:
        return _parse_xlsx(content_bytes), "xlsx"
    if normalized.startswith("text/"):
        return _parse_text(content_bytes), "text"
    if normalized in {"application/pdf"}:
        return _parse_binary(content_bytes), "pdf"
    if normalized in {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }:
        return _parse_binary(content_bytes), "office"
    return _parse_binary(content_bytes), "binary"
